# -*- coding: utf-8 -*-
"""定例レポートに「実CPA」「7日積み上げ」シートを追加する（2026-10-28 ユーザー要望）。
使い方: python3 scripts/extra_tabs.py <rep窓pkl> <report.xlsx>
実CPA = Meta広告費 ÷ Shopify販売数（個数）。分岐CPA = 売価 − 原価（バリアント加重は窓の平均原価）。
7日積み上げ = 直近7日の日次利益（実売−原価−広告）を1日目から積み上げ、前日比の増減を並べる。全店と商品別。"""
import sys, csv, pickle, collections, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
R = pickle.load(open(sys.argv[1], 'rb')); OUT = sys.argv[2]
B = 'data/daily/'
COST, PRICE, MAP = R['COST'], dict(R['PRICE']), R['MAP']
S = collections.defaultdict(lambda: collections.defaultdict(lambda: [0, 0]))
for r in csv.DictReader(open(B + 'daily_sales.csv', encoding='utf-8')):
    S[r['name']][r['date']][0] += int(r['gross']); S[r['name']][r['date']][1] += int(r['disc'])
AD = collections.defaultdict(lambda: collections.defaultdict(float))
for r in csv.DictReader(open(B + 'daily_ad.csv', encoding='utf-8')):
    AD[r['campaign']][r['date']] += float(r['cost'])
VG = collections.defaultdict(lambda: collections.defaultdict(lambda: collections.defaultdict(int)))
for r in csv.DictReader(open(B + 'daily_variant.csv', encoding='utf-8')):
    VG[r['name']][r['date']][r['cost_group']] += int(r['gross'])
VARC = {'ナノガラス脱毛パッド': {'白緑': 757, '黒桃': 740}, '壁掛けディスペンサー': {'3本': 2604, '2本': 2114}, '高見えレザーヘッドレストフック': {'黒': 1275, '茶灰': 1344}}
GP = {'ナノガラス脱毛パッド': {'白緑': 3980, '黒桃': 3980}, '高見えレザーヘッドレストフック': {'黒': 3980, '茶灰': 3980}, '壁掛けディスペンサー': {'3本': 6980, '2本': 5980}}
PRICE['優先配送'] = 790
def qc(n, d):
    g = S[n][d][0]
    if not g: return 0, 0
    if n in VARC:
        q = c = 0
        for grp, gg in VG[n][d].items():
            x = gg // GP[n][grp]; uc = 732 if (n == '高見えレザーヘッドレストフック' and d >= '2026-10-07') else VARC[n][grp]; q += x; c += x * uc
        return q, c
    p = PRICE.get(n)
    if not p or g % p: return None, None
    return g // p, g // p * COST.get(n, 0)
D1, D3, D7, D30 = R['D1'], R['D3'], R['D7'], R['D30']
ALLD = sorted({d for v in S.values() for d in v if d <= D1[0]})
D14 = ALLD[-14:]
def win(n, days):
    s = sum(S[n][d][0] - S[n][d][1] for d in days if d in S[n]); q = 0; c = 0; bad = False
    for d in days:
        if d not in S[n]: continue
        x, y = qc(n, d)
        if x is None: bad = True; continue
        q += x; c += y
    a = sum(AD[MAP.get(n, n)].get(d, 0) for d in days)
    return s, q, c, a, bad
wb = load_workbook(OUT)
TH = Font(name='Arial', bold=True, color='FFFFFF', size=10); TD = Font(name='Arial', size=10); NEG = Font(name='Arial', size=10, color='CC0000')
HEAD = PatternFill('solid', fgColor='305496'); thin = Border(*[Side(style='thin', color='CCCCCC')] * 4)
def sheet(title, note, cols, rows, widths, fmt, pos=None):
    if title in wb.sheetnames: del wb[title]
    ws = wb.create_sheet(title, pos) if pos is not None else wb.create_sheet(title); r = 1
    for ln in note: ws.cell(r, 1, ln).font = Font(name='Arial', size=10, bold=(r == 1)); r += 1
    r += 1
    for c, v in enumerate(cols, 1):
        x = ws.cell(r, c, v); x.font = TH; x.fill = HEAD; x.border = thin; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 30; r += 1
    for row in rows:
        for c, v in enumerate(row, 1):
            x = ws.cell(r, c, v); x.border = thin; x.font = NEG if isinstance(v, (int, float)) and v < 0 else TD
            if c in fmt: x.number_format = fmt[c]
        r += 1
    for i, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = ws.cell(len(note) + 3, 2)
# ---- 実CPA ----
rows = []
names = [n for n in S if n and n != '優先配送' and any(AD[MAP.get(n, n)].get(d, 0) for d in D30)]
for n in names:
    s7, q7, c7, a7, bad = win(n, D7)
    be = None
    if q7: be = round((s7 - c7) / q7)  # 1個あたり粗利（値引後・窓の加重原価）= 分岐CPA
    row = [n, be]
    for days in (D1, D3, D7, D14, D30):
        s, q, c, a, bad = win(n, days)
        row += [round(a / q) if q else None, q, round((s - c) / a, 2) if a else None]
    rows.append(row)
rows.sort(key=lambda r: -(r[9] or 0))
sheet('実CPA', [f'実CPA = Meta広告費 ÷ Shopify販売数（個数）。分岐CPA = 1個あたり粗利（値引後売上−原価、7日窓の実測）。実CPA < 分岐CPA なら黒字。窓: 前日 {D1[0][5:]} / 3日 {D3[0][5:]}〜 / 7日 {D7[0][5:]}〜 / 14日 {D14[0][5:]}〜 / 30日 {D30[0][5:]}〜',
  'Metaの購入数は注文単位（2個買われても1）かつ広告経由と判定した分だけなので、個数で見る実CPAの方が原価と直結する。倍率 = (売上−原価)÷広告'],
 ['商品', '分岐CPA', '前日 実CPA', '前日 個', '前日 倍率', '3日 実CPA', '3日 個', '3日 倍率', '7日 実CPA', '7日 個', '7日 倍率', '14日 実CPA', '14日 個', '14日 倍率', '30日 実CPA', '30日 個', '30日 倍率'],
 rows, [26, 9, 9, 6, 7, 9, 6, 7, 9, 6, 7, 9, 6, 7, 9, 6, 7], {2: '#,##0', 3: '#,##0', 5: '0.00', 6: '#,##0', 8: '0.00', 9: '#,##0', 11: '0.00', 12: '#,##0', 14: '0.00', 15: '#,##0', 17: '0.00'}, pos=2)
# ---- 7日積み上げ ----
WD = ['月', '火', '水', '木', '金', '土', '日']
def prof_day(d):
    s = sum(v[d][0] - v[d][1] for v in S.values() if d in v); k = 0
    for n in S:
        if d in S[n] and n:
            x, y = qc(n, d); k += (y or 0)
    a = sum(AD[c].get(d, 0) for c in AD)
    return s, k, a, s - k - a
rows = []; cum = 0; prev = None
for i, d in enumerate(D7, 1):
    s, k, a, p = prof_day(d); cum += p
    rows.append([f'{i}日目 {d[5:]}({WD[datetime.date.fromisoformat(d).weekday()]})', s, k, a, p, (p - prev) if prev is not None else None, cum])
    prev = p
hdr = ['日', '実売', '原価', '広告費', '日次利益', '前日比', '積み上げ利益']
# 商品別
prow = []
for n in names:
    cum = 0; prev = None; row = [n]
    for d in D7:
        s = (S[n][d][0] - S[n][d][1]) if d in S[n] else 0
        x, y = qc(n, d) if d in S[n] else (0, 0)
        a = AD[MAP.get(n, n)].get(d, 0); p = s - (y or 0) - a; cum += p
        row += [round(p), round(cum)]
    prow.append(row)
prow.sort(key=lambda r: -r[-1])
pcols = ['商品'] + [x for i, d in enumerate(D7, 1) for x in (f'{i}日目 利益', f'{i}日目 積み上げ')]
sheet('7日積み上げ', [f'直近7日（{D7[0][5:]}〜{D7[-1][5:]}）の日次利益を1日目から積み上げ。利益 = 実売（値引後）− 原価 − Meta広告費。前日比 = その日の利益 − 前日の利益', '下段は商品別（広告費はキャンペーン別、カタログ全部は全店のみに含む）'],
 hdr, rows, [18, 11, 10, 11, 11, 11, 13], {2: '#,##0', 3: '#,##0', 4: '#,##0', 5: '#,##0', 6: '#,##0', 7: '#,##0'}, pos=3)
ws = wb['7日積み上げ']; r = ws.max_row + 2
for c, v in enumerate(pcols, 1):
    x = ws.cell(r, c, v); x.font = TH; x.fill = HEAD; x.border = thin; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[r].height = 30; r += 1
for row in prow:
    for c, v in enumerate(row, 1):
        x = ws.cell(r, c, v); x.border = thin; x.font = NEG if isinstance(v, (int, float)) and v < 0 else TD
        if c > 1: x.number_format = '#,##0'
    r += 1
for i in range(2, len(pcols) + 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = 11
wb.save(OUT); print('実CPA / 7日積み上げ sheets added:', OUT)

# ---- 実CPA 週次推移（絶対評価つき）----
def wk(d): x = datetime.date.fromisoformat(d); return (x - datetime.timedelta(days=x.weekday())).isoformat()
weeks = sorted({wk(d) for d in ALLD if d >= (datetime.date.fromisoformat(D1[0]) - datetime.timedelta(days=62)).isoformat()})
rows = []
for n in names:
    row = [n]; be7 = None
    s7, q7, c7, a7, bad = win(n, D7); be = round((s7 - c7) / q7) if q7 else None
    row.append(be); series = []
    for w in weeks:
        ds = [d for d in ALLD if wk(d) == w]
        s, q, c, a, bad = win(n, ds)
        v = round(a / q) if (q and a) else None; series.append(v); row.append(v)
    full = [v for v in series[:-1] if v]  # 直近の週（進行中）を除いた確定週
    last_full = full[-1] if full else None
    margin = (be - last_full) if (be and last_full) else None
    margin_pct = round(margin / be * 100) if margin is not None else None
    hist_min = min(full) if full else None
    vs_best = round((last_full / hist_min - 1) * 100) if (last_full and hist_min) else None
    if margin_pct is None: verdict = '—'
    elif margin_pct >= 40: verdict = '◎ 余裕大'
    elif margin_pct >= 25: verdict = '○ 余裕あり'
    elif margin_pct >= 10: verdict = '△ 薄い'
    else: verdict = '✕ 分岐割れ'
    trend = '—'
    if len(full) >= 3:
        a3 = sum(full[-3:]) / 3; b3 = sum(full[-6:-3]) / 3 if len(full) >= 6 else (sum(full[:-3]) / len(full[:-3]) if full[:-3] else None)
        if b3: trend = '↑悪化' if a3 > b3 * 1.15 else ('↓改善' if a3 < b3 * 0.85 else '→横ばい')
    rows.append(row + [last_full, margin, margin_pct, hist_min, vs_best, trend, verdict])
rows.sort(key=lambda r: -(r[-5] or -1))
cols = ['商品', '分岐CPA(7日実測)'] + [f'{w[5:]}週' for w in weeks] + ['直近確定週 実CPA', '余裕額(分岐−実CPA)', '余裕率%', '過去最良週', '最良比%', '直近3週の傾向', '絶対評価']
sheet('実CPA週次', [f'週別 実CPA（広告÷Shopify個数・月曜始まり）。最後の週は進行中（{D1[0][5:]}まで）。絶対評価は直近の確定週の実CPAを分岐CPAと比べたもの（余裕率 = (分岐−実CPA)÷分岐）',
  '◎ 余裕40%以上 ／ ○ 25〜40% ／ △ 10〜25% ／ ✕ 分岐割れ。傾向 = 直近3確定週の平均 vs その前3週（±15%で判定）。「過去より悪い」と「分岐より余裕がある」は別の軸なので両方を併記'],
 cols, rows, [26, 10] + [8] * len(weeks) + [10, 11, 8, 9, 8, 10, 10], {c: '#,##0' for c in range(2, len(cols) + 1)}, pos=3)
wb.save(OUT); print('実CPA週次 sheet added')
